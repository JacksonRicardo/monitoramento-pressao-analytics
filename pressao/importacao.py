"""Importação de medições a partir de planilhas no modelo original."""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook

from .models import Medicao, Periodo, periodo_por_hora


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return texto.strip().lower()


# Prefixos aceitos para identificar cada coluna do cabeçalho
CABECALHOS = {
    "data": ("data",),
    "hora": ("hora", "horario"),
    "sistolica": ("sistolica", "pas", "maxima"),
    "diastolica": ("diastolica", "pad", "minima"),
    "pulsacao": ("pulsacao", "pulso", "bpm", "frequencia"),
    "periodo": ("periodo", "turno"),
    "nota_profissional": ("nota", "nota do profissional"),
    "observacoes": ("observ", "obs"),
}

PERIODOS = {_normalizar(v): v for v, _ in Periodo.choices}


@dataclass
class ResultadoImportacao:
    criadas: int = 0
    ignoradas: int = 0
    erros: list[str] = field(default_factory=list)


def _para_data(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(valor))
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except ValueError:
            continue
    raise ValueError(f"data inválida '{valor}'")


def _para_hora(valor) -> time | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    if isinstance(valor, timedelta):
        minutos = int(valor.total_seconds() // 60) % (24 * 60)
        return time(minutos // 60, minutos % 60)
    if isinstance(valor, (int, float)):
        minutos = round(float(valor) % 1 * 24 * 60) % (24 * 60)
        return time(minutos // 60, minutos % 60)
    for formato in ("%H:%M", "%H:%M:%S", "%Hh%M", "%Hh"):
        try:
            return datetime.strptime(str(valor).strip(), formato).time()
        except ValueError:
            continue
    raise ValueError(f"hora inválida '{valor}'")


def _para_int(valor) -> int | None:
    if valor is None or valor == "":
        return None
    try:
        return int(round(float(str(valor).replace(",", "."))))
    except ValueError as exc:
        raise ValueError(f"número inválido '{valor}'") from exc


def _mapear_cabecalho(linha) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for idx, celula in enumerate(linha):
        nome = _normalizar(celula)
        if not nome:
            continue
        for chave, prefixos in CABECALHOS.items():
            if chave not in mapa and any(nome.startswith(p) for p in prefixos):
                mapa[chave] = idx
                break
    return mapa


def importar_planilha(arquivo, usuario, ignorar_duplicadas: bool = True) -> ResultadoImportacao:
    resultado = ResultadoImportacao()
    try:
        wb = load_workbook(arquivo, data_only=True, read_only=True)
    except Exception:  # arquivo corrompido ou formato diferente
        resultado.erros.append("Não foi possível ler o arquivo. Verifique se é uma planilha .xlsx válida.")
        return resultado

    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()

    # Procura a linha de cabeçalho nas primeiras 15 linhas
    mapa, inicio = {}, 0
    for i, linha in enumerate(linhas[:15]):
        candidato = _mapear_cabecalho(linha)
        if {"data", "sistolica", "diastolica"} <= candidato.keys():
            mapa, inicio = candidato, i + 1
            break
    if not mapa:
        resultado.erros.append(
            "Cabeçalho não encontrado. A planilha precisa ter as colunas Data, Hora, Sistólica e Diastólica."
        )
        return resultado

    existentes = set()
    if ignorar_duplicadas:
        existentes = set(Medicao.objects.filter(usuario=usuario).values_list("data", "hora"))

    def pegar(linha, chave):
        idx = mapa.get(chave)
        return linha[idx] if idx is not None and idx < len(linha) else None

    novas = []
    for numero, linha in enumerate(linhas[inicio:], start=inicio + 1):
        if not linha or all(v in (None, "") for v in linha):
            continue
        primeira = _normalizar(linha[0])
        # Fim da área de dados (ex.: bloco "ESTATÍSTICAS" da planilha original)
        if primeira.startswith(("estatistica", "media", "total")):
            break
        try:
            data_m = _para_data(pegar(linha, "data"))
            sis = _para_int(pegar(linha, "sistolica"))
            dia = _para_int(pegar(linha, "diastolica"))
            if data_m is None or sis is None or dia is None:
                continue
            hora_m = _para_hora(pegar(linha, "hora")) or time(0, 0)
            pulso = _para_int(pegar(linha, "pulsacao"))
            periodo = PERIODOS.get(_normalizar(pegar(linha, "periodo")), "") or periodo_por_hora(hora_m)

            if not (50 <= sis <= 300) or not (30 <= dia <= 200):
                raise ValueError(f"valores fora do intervalo aceito ({sis}/{dia})")
            if dia >= sis:
                raise ValueError("diastólica maior ou igual à sistólica")
            if pulso is not None and not (25 <= pulso <= 250):
                raise ValueError(f"pulsação fora do intervalo aceito ({pulso})")
        except ValueError as exc:
            resultado.erros.append(f"Linha {numero}: {exc}")
            continue

        if (data_m, hora_m) in existentes:
            resultado.ignoradas += 1
            continue
        existentes.add((data_m, hora_m))
        novas.append(
            Medicao(
                usuario=usuario,
                data=data_m,
                hora=hora_m,
                sistolica=sis,
                diastolica=dia,
                pulsacao=pulso,
                periodo=periodo,
                nota_profissional=str(pegar(linha, "nota_profissional") or "")[:255],
                observacoes=str(pegar(linha, "observacoes") or "")[:255],
            )
        )
    Medicao.objects.bulk_create(novas)
    resultado.criadas = len(novas)
    return resultado

