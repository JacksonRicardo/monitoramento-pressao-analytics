"""Geração da planilha Excel customizada (baseada no modelo original)."""
from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Periodo

INSTRUCOES = [
    ("COMO USAR ESTA PLANILHA", True),
    ("", False),
    ("1. REGISTRAR MEDIÇÕES:", True),
    ("   • Data: formato DD/MM/AAAA", False),
    ("   • Hora: formato HH:MM", False),
    ("   • Sistólica: pressão máxima (ex: 120)", False),
    ("   • Diastólica: pressão mínima (ex: 80)", False),
    ("   • Pulsação: batimentos por minuto", False),
    ("   • Período: escolha entre Manhã, Tarde ou Noite", False),
    ("", False),
    ("2. REGISTRO, SEM DIAGNÓSTICO AUTOMÁTICO:", True),
    ("   • Nota do profissional: anotação opcional feita com orientação profissional.", False),
    ("   • As faixas de referência são apenas informativas e não substituem avaliação médica.", False),
    ("   • Médias e gráfico resumem os números; não definem uma meta clínica.", False),
    ("   • Máximos e mínimos são separados e podem vir de medições diferentes.", False),
    ("   • Se houver sintomas importantes, procure atendimento; não espere a planilha.", False),
    ("", False),
    ("3. DICAS IMPORTANTES:", True),
    ("   • Combine horários e frequência de registro com seu profissional.", False),
    ("   • Fique em repouso 5 minutos antes de medir", False),
    ("   • Não fale durante a medição", False),
    ("   • Anote observações relevantes", False),
    ("   • Leve este controle em suas consultas médicas", False),
    ("", False),
    ("Consulte sempre seu médico para orientações sobre sua pressão arterial.", True),
]


def _clarear(hex_cor: str, fator: float = 0.85) -> str:
    r, g, b = (int(hex_cor[i : i + 2], 16) for i in (0, 2, 4))
    r, g, b = (int(c + (255 - c) * fator) for c in (r, g, b))
    return f"{r:02X}{g:02X}{b:02X}"


def _colunas(opcoes: dict) -> list[tuple[str, str, int]]:
    """(chave, título, largura) na ordem da planilha original."""
    cols = [
        ("data", "Data", 13),
        ("hora", "Hora", 9),
        ("sistolica", "Sistólica (máxima)", 13),
        ("diastolica", "Diastólica (mínima)", 13),
    ]
    extras = {
        "pulsacao": ("pulsacao", "Pulsação (bpm)", 12),
        "periodo": ("periodo", "Período", 11),
        "pam": ("pam", "PAM", 8),
        "classificacao": ("classificacao", "Faixa de referência", 18),
        "nota_profissional": ("nota_profissional", "Nota do profissional", 28),
        "observacoes": ("observacoes", "Observações", 28),
    }
    for chave in ("pulsacao", "periodo", "pam", "classificacao", "nota_profissional", "observacoes"):
        if chave in opcoes["colunas"]:
            cols.append(extras[chave])
    return cols


def _valor(medicao, chave):
    if chave == "classificacao":
        return medicao.classificacao["rotulo"]
    if chave == "pam":
        return medicao.pressao_media_arterial
    valor = getattr(medicao, chave)
    return valor if valor not in (None, "") else None


def gerar_excel(medicoes, analise: dict, opcoes: dict, usuario) -> bytes:
    cor = opcoes["cor"]
    cor_clara = _clarear(cor)
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    fonte_titulo = Font(bold=True, size=16, color="FFFFFF")
    fill_titulo = PatternFill("solid", fgColor=cor)
    fill_cabecalho = PatternFill("solid", fgColor=cor_clara)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Registro de Pressão"

    cols = _colunas(opcoes)
    ncols = len(cols)
    ultima_col = get_column_letter(ncols)

    # ---- Cabeçalho ----------------------------------------------------------
    ws.merge_cells(f"A1:{ultima_col}1")
    ws["A1"] = opcoes["titulo"].upper()
    ws["A1"].font = fonte_titulo
    ws["A1"].fill = fill_titulo
    ws["A1"].alignment = centro
    ws.row_dimensions[1].height = 32

    nome = usuario.get_full_name() or usuario.username
    intervalo = ""
    if not analise["vazio"]:
        intervalo = f" | Período: {analise['inicio']:%d/%m/%Y} a {analise['fim']:%d/%m/%Y}"
    ws.merge_cells(f"A2:{ultima_col}2")
    ws["A2"] = (
        f"Paciente: {nome}{intervalo} | Gerado em {timezone.localtime():%d/%m/%Y %H:%M}"
    )
    ws["A2"].font = Font(italic=True, size=10, color="595959")
    ws["A2"].alignment = Alignment(horizontal="center")

    linha_cab = 3
    for i, (_, titulo, largura) in enumerate(cols, start=1):
        c = ws.cell(row=linha_cab, column=i, value=titulo)
        c.font = Font(bold=True, color="000000")
        c.fill = fill_cabecalho
        c.alignment = centro
        c.border = borda
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha_cab].height = 30

    # ---- Dados --------------------------------------------------------------
    linha = linha_cab + 1
    for m in medicoes:
        for i, (chave, _, _) in enumerate(cols, start=1):
            c = ws.cell(row=linha, column=i, value=_valor(m, chave))
            c.border = borda
            if chave == "data":
                c.number_format = "DD/MM/YYYY"
                c.alignment = Alignment(horizontal="center")
            elif chave == "hora":
                c.number_format = "HH:MM"
                c.alignment = Alignment(horizontal="center")
            elif chave in ("sistolica", "diastolica", "pulsacao", "pam", "periodo"):
                c.alignment = Alignment(horizontal="center")
            elif chave == "classificacao":
                c.alignment = Alignment(horizontal="center")
                c.font = Font(color=m.classificacao["hex"].lstrip("#"), bold=True)
        linha += 1

    if not medicoes:
        # Modelo em branco: 20 linhas prontas para preenchimento, como na planilha original
        fill_preencher = PatternFill("solid", fgColor="FFF9DB")
        for _ in range(20):
            for i, (chave, _, _) in enumerate(cols, start=1):
                c = ws.cell(row=linha, column=i)
                c.border = borda
                c.fill = fill_preencher
                c.alignment = Alignment(horizontal="center")
                if chave == "data":
                    c.number_format = "DD/MM/YYYY"
                elif chave == "hora":
                    c.number_format = "HH:MM"
            linha += 1

    primeira, ultima = linha_cab + 1, max(linha - 1, linha_cab + 1)
    idx = {chave: get_column_letter(i) for i, (chave, _, _) in enumerate(cols, start=1)}

    ws.freeze_panes = ws.cell(row=linha_cab + 1, column=1)
    ws.auto_filter.ref = f"A{linha_cab}:{ultima_col}{ultima}"

    # Lista suspensa no Período (como na planilha original), inclusive para novas linhas
    if "periodo" in idx:
        dv = DataValidation(
            type="list", formula1='"' + ",".join(v for v, _ in Periodo.choices) + '"', allow_blank=True
        )
        dv.error = "Escolha Manhã, Tarde ou Noite"
        dv.prompt = "Selecione o período"
        ws.add_data_validation(dv)
        dv.add(f"{idx['periodo']}{primeira}:{idx['periodo']}{ultima + 200}")

    # Destaque condicional (continua funcionando se o usuário editar a planilha)
    if opcoes["destacar"]:
        s, d = idx["sistolica"], idx["diastolica"]
        faixa = f"A{primeira}:{ultima_col}{ultima}"
        ws.conditional_formatting.add(
            faixa,
            FormulaRule(
                formula=[f"OR(${s}{primeira}>=180,${d}{primeira}>=110)"],
                fill=PatternFill("solid", fgColor="F5B7B1"),
                font=Font(bold=True, color="78281F"),
                stopIfTrue=True,
            ),
        )
        ws.conditional_formatting.add(
            faixa,
            FormulaRule(
                formula=[f"OR(${s}{primeira}>=140,${d}{primeira}>=90)"],
                fill=PatternFill("solid", fgColor="FDEBD0"),
            ),
        )

    # ---- Estatísticas (fórmulas vivas, como na planilha original) ----------
    if opcoes["estatisticas"]:
        s, d = idx["sistolica"], idx["diastolica"]
        rs, rd = f"{s}{primeira}:{s}{ultima}", f"{d}{primeira}:{d}{ultima}"
        base = ultima + 3
        ws.merge_cells(f"A{base}:{ultima_col}{base}")
        ws[f"A{base}"] = "ESTATÍSTICAS"
        ws[f"A{base}"].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f"A{base}"].fill = fill_titulo
        ws[f"A{base}"].alignment = centro

        estat = [
            ("Média sistólica", f'=IF(COUNT({rs})>0,ROUND(AVERAGE({rs}),0),"")'),
            ("Média diastólica", f'=IF(COUNT({rd})>0,ROUND(AVERAGE({rd}),0),"")'),
        ]
        if "pulsacao" in idx:
            p = idx["pulsacao"]
            rp = f"{p}{primeira}:{p}{ultima}"
            estat.append(("Média de pulsação", f'=IF(COUNT({rp})>0,ROUND(AVERAGE({rp}),0),"")'))
        estat += [
            ("Mediana sistólica / diastólica", f'=IF(COUNT({rs})>0,MEDIAN({rs})&" / "&MEDIAN({rd}),"")'),
            ("Desvio padrão sistólica", f'=IF(COUNT({rs})>1,ROUND(STDEV({rs}),1),"")'),
            ("Desvio padrão diastólica", f'=IF(COUNT({rd})>1,ROUND(STDEV({rd}),1),"")'),
            ("Máximos (sist. / diast.)", f'=IF(COUNT({rs})>0,MAX({rs})&" / "&MAX({rd}),"")'),
            ("Mínimos (sist. / diast.)", f'=IF(COUNT({rs})>0,MIN({rs})&" / "&MIN({rd}),"")'),
            ("Total de medições", f"=COUNT({rs})"),
            ("Leituras ≥ 140/90", f'=SUMPRODUCT(--((({rs})>=140)+(({rd})>=90)>0))'),
        ]
        for i, (rotulo, formula) in enumerate(estat, start=base + 1):
            ws.merge_cells(f"A{i}:B{i}")
            ws[f"A{i}"] = rotulo
            ws[f"A{i}"].font = Font(bold=True)
            ws[f"A{i}"].fill = fill_cabecalho
            ws[f"A{i}"].border = borda
            ws[f"C{i}"] = formula
            ws[f"C{i}"].alignment = Alignment(horizontal="center")
            ws[f"C{i}"].border = borda
            ws[f"C{i}"].font = Font(bold=True)

    # Impressão
    ws.print_title_rows = f"{linha_cab}:{linha_cab}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    # ---- Gráfico de evolução ------------------------------------------------
    if opcoes["grafico"] and medicoes:
        wg = wb.create_sheet("Gráfico de Evolução")
        grafico = LineChart()
        grafico.title = "Evolução da Pressão Arterial"
        grafico.y_axis.title = "mmHg"
        grafico.x_axis.title = "Data"
        grafico.x_axis.number_format = "DD/MM"
        grafico.height, grafico.width = 11, 26
        col_s = cols.index(next(c for c in cols if c[0] == "sistolica")) + 1
        col_d = col_s + 1
        dados = Reference(ws, min_col=col_s, max_col=col_d, min_row=linha_cab, max_row=ultima)
        grafico.add_data(dados, titles_from_data=True)
        grafico.set_categories(Reference(ws, min_col=1, min_row=primeira, max_row=ultima))
        for serie, cor_linha in zip(grafico.series, ("C0392B", "2E86C1")):
            serie.graphicalProperties.line.solidFill = cor_linha
            serie.graphicalProperties.line.width = 22000
            serie.smooth = True
        if "pulsacao" in idx:
            col_p = cols.index(next(c for c in cols if c[0] == "pulsacao")) + 1
            grafico.add_data(
                Reference(ws, min_col=col_p, min_row=linha_cab, max_row=ultima), titles_from_data=True
            )
            grafico.series[-1].graphicalProperties.line.solidFill = "27AE60"
            grafico.series[-1].graphicalProperties.line.dashStyle = "dash"
        wg.add_chart(grafico, "A1")

    # ---- Análise ------------------------------------------------------------
    if opcoes["analise"] and not analise["vazio"]:
        wa = wb.create_sheet("Análise")
        wa.column_dimensions["A"].width = 24
        for letra in "BCDE":
            wa.column_dimensions[letra].width = 16

        def titulo_secao(linha_t, texto):
            wa.merge_cells(f"A{linha_t}:E{linha_t}")
            wa[f"A{linha_t}"] = texto
            wa[f"A{linha_t}"].font = Font(bold=True, color="FFFFFF", size=12)
            wa[f"A{linha_t}"].fill = fill_titulo
            wa[f"A{linha_t}"].alignment = centro

        def cabecalho(linha_c, valores):
            for i, v in enumerate(valores, start=1):
                c = wa.cell(row=linha_c, column=i, value=v)
                c.font = Font(bold=True)
                c.fill = fill_cabecalho
                c.border = borda
                c.alignment = centro

        def linha_dados(linha_d, valores):
            for i, v in enumerate(valores, start=1):
                c = wa.cell(row=linha_d, column=i, value=v)
                c.border = borda
                if i > 1:
                    c.alignment = Alignment(horizontal="center")

        titulo_secao(1, "MÉDIAS POR PERÍODO DO DIA")
        cabecalho(2, ["Período", "Medições", "Sistólica", "Diastólica", "Pulsação"])
        r = 3
        for p in analise["por_periodo"]:
            linha_dados(r, [p["periodo"], p["n"], p["sistolica"], p["diastolica"], p["pulsacao"]])
            r += 1
        bar = BarChart()
        bar.title = "Média por período"
        bar.add_data(Reference(wa, min_col=3, max_col=4, min_row=2, max_row=r - 1), titles_from_data=True)
        bar.set_categories(Reference(wa, min_col=1, min_row=3, max_row=r - 1))
        bar.height, bar.width = 7, 12
        wa.add_chart(bar, "G1")

        r += 1
        titulo_secao(r, "DISTRIBUIÇÃO POR FAIXA DE REFERÊNCIA (informativa)")
        cabecalho(r + 1, ["Faixa", "Medições", "%", "", ""])
        inicio_dist = r + 2
        r = inicio_dist
        for f in analise["distribuicao"]:
            linha_dados(r, [f["rotulo"], f["n"], f["pct"] / 100, None, None])
            wa.cell(row=r, column=3).number_format = "0.0%"
            wa.cell(row=r, column=1).font = Font(bold=True, color=f["hex"].lstrip("#"))
            r += 1
        pizza = PieChart()
        pizza.title = "Faixas de referência"
        pizza.add_data(Reference(wa, min_col=2, min_row=inicio_dist, max_row=r - 1))
        pizza.set_categories(Reference(wa, min_col=1, min_row=inicio_dist, max_row=r - 1))
        pizza.dataLabels = DataLabelList()
        pizza.dataLabels.showPercent = True
        pizza.height, pizza.width = 8, 12
        wa.add_chart(pizza, "G16")

        r += 1
        titulo_secao(r, "MÉDIAS POR DIA DA SEMANA")
        cabecalho(r + 1, ["Dia", "Medições", "Sistólica", "Diastólica", ""])
        r += 2
        for dsem in analise["por_dia_semana"]:
            linha_dados(r, [dsem["dia"], dsem["n"], dsem["sistolica"], dsem["diastolica"], None])
            r += 1

        r += 1
        titulo_secao(r, "INDICADORES")
        r += 1
        indicadores = [
            ("Total de medições", analise["total"]),
            ("Dias com registro", analise["dias_com_registro"]),
            ("Regularidade de registro", f"{analise['regularidade']}%"),
            ("PAM média", analise["pam"]["media"]),
            ("Pressão de pulso média", analise["pressao_pulso"]["media"]),
            ("Variabilidade sistólica (CV)", f"{analise['sistolica']['cv']}%"),
            ("Variabilidade diastólica (CV)", f"{analise['diastolica']['cv']}%"),
        ]
        if analise.get("tendencia"):
            t = analise["tendencia"]
            indicadores += [
                ("Tendência sistólica", f"{t['sistolica']:+} mmHg/semana ({t['sistolica_rotulo']})"),
                ("Tendência diastólica", f"{t['diastolica']:+} mmHg/semana ({t['diastolica_rotulo']})"),
            ]
        if analise.get("meta"):
            mt = analise["meta"]
            indicadores.append(
                (f"Dentro da meta ({mt['sistolica']}/{mt['diastolica']})", f"{mt['pct']}% ({mt['dentro']})")
            )
        for rotulo, valor in indicadores:
            wa.merge_cells(f"B{r}:C{r}")
            linha_dados(r, [rotulo, valor])
            wa.cell(row=r, column=1).font = Font(bold=True)
            r += 1

    # ---- Instruções ---------------------------------------------------------
    if opcoes["instrucoes"]:
        wi = wb.create_sheet("Instruções")
        wi.column_dimensions["A"].width = 95
        for i, (texto, negrito) in enumerate(INSTRUCOES, start=1):
            wi.cell(row=i, column=1, value=texto).font = Font(
                bold=negrito, size=14 if i == 1 else 11, color=cor if negrito else "000000"
            )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
